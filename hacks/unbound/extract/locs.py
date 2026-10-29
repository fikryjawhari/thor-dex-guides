"""Wiki encounters and raid dens, supplemented by the public 2.1.1.1 location workbook.

Floors and dens are sections, not separate tracker slots. P.w retains source method, time and rate;
the sixth field hides rows rendered through AR.es. Wild levels are deliberately absent.
"""
from collections import Counter
from pathlib import Path
import re

import openpyxl

from common import H, RD, WIKI, report, sid
from wiki import key, text


ORDER = ['Shadow Base', 'Frozen Heights', 'Route 1', 'Bellin Town', 'Icicle Cave', 'Route 2', 'Dresco Town',
         'Route 3', 'Flower Paradise', 'Grim Woods', 'Route 4', 'Cinder Volcano', 'Route 5', 'Crater Town',
         'KBT Expressway', 'Route 6', 'Valley Cave', 'Route 7', 'Frost Mountain', 'Route 8', 'Blizzard City',
         'Frozen Forest', 'Route 9', 'Fallshore City', 'Route 11', 'Epidimy Town', 'Route 10',
         'Tehl Town / Underground Pass', 'Thundercap Mountain', 'Tarmigan Town', 'Tarmigan Mansion',
         'Route 12', 'Dehara City', 'Ruins of Void', 'Cinder Volcano Depths', 'Route 13', 'Route 14',
         'Auburn Waterway', 'Lost Tunnel', 'Great Desert', 'Gurun Town', 'Vivill Woods', 'Vivill Town',
         'Vivill Warehouse', 'Route 15', 'Route 16', 'Antisis City', 'Antisis Port', 'Antisis Sewers',
         'Ferrothorn Turf', 'Icy Hole', 'Route 17', 'Seaport City', 'Seaport Warehouses', 'S.S. Marine',
         'Polder Town', 'Cootes Bog', 'Magnolia Town', 'Magnolia Fields', 'Crystal Peak', 'Redwood Village',
         'Redwood Forest', 'Route 18', 'Victory Road', 'Battle Frontier', 'Cube Corp.', 'Cliff Cave']

AREA_ALIASES = {'Ice Hole': 'Icy Hole', 'Thundercap Mt.': 'Thundercap Mountain', 'Antisis Sewers': 'Antisis Sewers',
                'Frozen Height': 'Frozen Heights', 'Route 9 + Autl Woods': 'Route 9',
                'Tehl Town': 'Tehl Town / Underground Pass', 'RoV (Underwater)': 'Ruins of Void',
                'Mision HQ': 'Fallshore City', 'Mission HQ': 'Fallshore City', 'Safari Zone': 'Polder Town',
                'Magnolia Café': 'Magnolia Town', 'Battleground': 'Seaport City',
                'Rift Cave': 'Dehara City', 'Tomb of Borrius': 'Great Desert', 'Dream Lab': 'Tarmigan Town',
                'Cave of Being': 'Route 14', 'Fullmoon Island': 'Tarmigan Town', 'Newmoon Island': 'Tarmigan Town',
                'Distortion World': 'Vivill Woods', 'Pokémon Center': 'Bellin Town'}

NAME_FIX = {'Klank': 'Klang', 'Whirlpide': 'Whirlipede', 'Amoongus': 'Amoonguss', 'Spheal Egg': 'Spheal',
            'Flabebe': 'Flabébé', 'Flabébé': 'Flabébé', 'Nidoran♂️': 'Nidoran♂', 'Nidoran♀️': 'Nidoran♀'}
GROUPS = {'One of the Borrius Starters': ['Beldum', 'Gible', 'Larvitar'],
          'Unova Starters': ['Snivy', 'Tepig', 'Oshawott'], 'Sinnoh Starters': ['Turtwig', 'Chimchar', 'Piplup'],
          'Alola Starters': ['Rowlet', 'Litten', 'Popplio'], 'Kanto Starters': ['Bulbasaur', 'Charmander', 'Squirtle'],
          'Kalos Starters': ['Chespin', 'Fennekin']}
NUMERIC = re.compile(r'\d+(?:\.\d+)?%$')


def species_names(name):
    if name == 'Jellicent (M, F)':
        return ['Jellicent', 'Jellicent (F)']
    if name == 'Basculin (Red and Blue)':
        return ['Basculin (Red Striped)', 'Basculin (Blue Striped)']
    return [NAME_FIX.get(name, name)]


def species_ids(name):
    out = list(dict.fromkeys(sid(value) for value in species_names(name)))
    if None in out:
        raise ValueError('Unresolved location species: %s' % name)
    return out


def area_name(value):
    value = str(value).strip()
    if value in AREA_ALIASES:
        return AREA_ALIASES[value]
    for name in sorted(ORDER, key=len, reverse=True):
        if key(value).startswith(key(name)):
            return name
    for alias in sorted(AREA_ALIASES, key=len, reverse=True):
        if key(value).startswith(key(alias)):
            return AREA_ALIASES[alias]
    raise ValueError('Unresolved location: %s' % value)


def workbook():
    return openpyxl.load_workbook(Path(H.src) / 'research-2026-10-29/location-guide.xlsx', data_only=True)


def doc_wild(book):
    """The two wide membership sheets: location columns, floor/method separators, colour-coded time."""
    records, labels = [], Counter()
    for sheet_name, header in [('Grass & Cave Encounters', 1), ('Surfing, Fishing, Rock Smash', 3)]:
        sheet = book[sheet_name]
        for column in range(1, sheet.max_column + 1):
            heading = sheet.cell(header, column).value
            if not heading:
                continue
            area = area_name(heading)
            method, floor = ('Grass/cave' if header == 1 else 'Surf'), ''
            for row in range(header + 1, sheet.max_row + 1):
                cell = sheet.cell(row, column)
                if not cell.value:
                    continue
                value = str(cell.value).strip()
                value = re.sub(r'\s+', ' ', value)
                candidate = re.sub(r'\s*\((?:Daily|Weekly|Specific days)\)', '', value, flags=re.I).strip()
                names = GROUPS.get(candidate, species_names(candidate))
                if all(sid(name) for name in names):
                    colour = cell.fill.fgColor
                    colour = colour.rgb if colour.type == 'rgb' else str(colour.indexed)
                    records.append((area, str(heading), floor, method, names, colour, cell.coordinate))
                elif value in ('Swarm', 'Special Encounter', 'Old Rod', 'Good Rod', 'Super Rod', 'Rock Smash',
                               'Headbutt', 'Surfing', 'Surf'):
                    method = {'Special Encounter': 'DexNav', 'Surfing': 'Surf'}.get(value, value)
                elif value == 'Fishing':
                    continue
                else:
                    labels[value] += 1
                    if re.search(r'\b[1-9]F\b|B[1-9]F|Flowers|Easy|Medium|Hard|Entrance|Underwater|Summit|Outside', value):
                        floor = value
                        if method in ('Swarm', 'DexNav'):
                            method = 'Grass/cave' if header == 1 else 'Surf'
    return records, labels


def doc_specials(book):
    records = []
    sheet = book['Gift & Static Encounters']
    for row in range(4, 53):
        method, place, group, requirement = (sheet.cell(row, column).value for column in (1, 3, 7, 8))
        value = sheet.cell(row, 5).value or sheet.cell(row, 6).value
        if not place:
            if row in (50, 51):
                place, method, group = 'Magnolia Café', 'Egg', ('Kalos Starters' if row == 50 else 'Froakie')
                requirement = 'Froakie post-game' if row == 51 else 'Free egg once a day'
            else:
                continue
        names = GROUPS.get(group)
        if names is None and (value or (group and sid(group))):
            value = str(value or group)
            clean = re.sub(r'\s*\((?:Daily|Weekly|Specific days)\)', '', value, flags=re.I)
            clean = re.sub(r'^(?:Shiny |Gigantamax )+', '', clean).removesuffix(' Egg')
            names = [NAME_FIX.get(name.strip(), name.strip()) for name in clean.split('/')]
        if not names:
            continue
        note = '; '.join(str(part) for part in (place, requirement, value or group) if part and part != 'None')
        records.append((area_name(place), str(method), names, note))
    sheet = book['LegendMythical & Ultra Beasts']
    for row in (3, 14, 25, 36, 47):
        for column in range(1, sheet.max_column, 2):
            value = sheet.cell(row, column).value
            if not value:
                continue
            names = {'Swords of Justice': ['Cobalion', 'Terrakion', 'Virizion']}.get(value)
            if names is None:
                names = re.split(r'\s*(?:\+|,|/)\s*', str(value))
            place = sheet.cell(row + 2, column + 1).value
            if not place:
                if value == 'Phione':
                    continue
                raise ValueError('Legendary without place: %s' % value)
            notes = [str(sheet.cell(offset, col).value) for offset in range(row + 2, min(row + 10, sheet.max_row + 1))
                     for col in (column, column + 1) if sheet.cell(offset, col).value]
            method = 'Gift' if 'Mystery Gift' in notes else 'Roaming' if 'Roaming' in notes else 'Static'
            records.append((area_name(place), method, names, '; '.join(notes)))
    sheet = book['In-game trades & Game Corner']
    for column in (1, 6, 11, 16, 21, 26):
        area = area_name(sheet.cell(4, column).value)
        for row in (6, 7):
            give, receive = sheet.cell(row, column).value, sheet.cell(row, column + 2).value
            if receive:
                records.append((area, 'Trade', [NAME_FIX.get(receive, receive)], 'Trade for %s' % give))
        records.append(('Dehara City', 'Game Corner', [sheet.cell(12, column).value],
                        '%s coins' % int(sheet.cell(11, column).value)))
    return records


def rhd_membership():
    out = {}
    for path in sorted((Path(RD.path) / 'locations').glob('*/index.html')):
        raw = path.read_text(encoding='utf-8')
        title = text(re.search(r'<h1\b[^>]*>(.*?)</h1>', raw, re.S).group(1))
        names = re.findall(r'<tr\b[^>]*class="[^"]*encounter-row[^>]*data-id="([^"]+)"', raw)
        out[title] = {species_id for name in names for species_id in species_ids(name)}
    return out


def membership_differences(data):
    by_area = {area['n']: {species['id'] for species in data['P'] for row in species['w'] if row[0] in area['l']}
               for area in data['AR']}
    differences = []
    for name, members in rhd_membership().items():
        try:
            area = area_name(name)
        except ValueError:
            differences.append((name, sorted(members), []))
            continue
        ours = by_area[area]
        if members != ours:
            differences.append((name, sorted(members - ours), sorted(ours - members)))
    return differences


def build(data):
    pages = {page.title: page for page in WIKI.locations()}
    if set(pages) != set(ORDER):
        raise ValueError('Location page inventory changed: %s' % (set(pages) ^ set(ORDER)))
    data['LOC'] = ['Starter'] + ORDER.copy()
    data['AR'] = [{'n': name, 'l': [index]} for index, name in enumerate(data['LOC'])]
    data['AR'][0]['st'] = 1
    by_area = {area['n']: area for area in data['AR']}
    by_id = {species['id']: species for species in data['P']}
    for species in data['P']:
        species['w'] = []
    counts = Counter()

    def note(area, value):
        if value and value not in by_area[area].get('note', ''):
            by_area[area]['note'] = (by_area[area].get('note', '') + ' ' + value).strip()

    def add(area, name, method, rate='', when='', floor='', section='', detail=''):
        for species_id in species_ids(name):
            species = by_id[species_id]
            row_note = '; '.join(part for part in (when, floor, detail) if part and part not in ('–', '-'))
            row = [by_area[area]['l'][0], method, rate, '', row_note]
            if section:
                row.append(1)
            if row in species['w']:
                continue
            species['w'].append(row)
            if section:
                sections = by_area[area].setdefault('es', [])
                label = section if not floor or section == floor else floor + ': ' + section
                box = next((box for box in sections if box['n'] == label), None)
                if box is None:
                    box = {'n': label, 'k': 'special', 'g': []}
                    sections.append(box)
                label = method + (' · ' + when if when and when not in ('–', '-') else '')
                group = next((group for group in box['g'] if group['m'] == label), None)
                if group is None:
                    group = {'m': label, 'rows': []}
                    box['g'].append(group)
                value = float(rate.rstrip('%')) if NUMERIC.fullmatch(rate) else 0
                group['rows'].append({'p': species_id, 'r': [value, value], 'note': detail})
            counts['rows'] += 1

    for name in GROUPS['One of the Borrius Starters']:
        add('Starter', name, 'Starter choice', '1 of 3', when='Shadow Base')
    for area in ORDER:
        for floor, name, method, when, rate in pages[area].encounters:
            if method == 'Starter':
                continue
            if NUMERIC.fullmatch(rate):
                counts['wiki wild'] += 1
                add(area, name, method, rate, when, floor, section=floor or 'Wild Pokémon')
            else:
                counts['wiki other'] += 1
                detail = 'Gigantamax' if 'Gigantamax' in name else 'Shiny' if 'Shiny' in name else ''
                add(area, name, method, rate if rate == '?%' else '', when, floor, detail=detail)
            if 'Special Encounter' in method:
                note(area, 'Gen 8: register a caught #810-890 species (not a regional form); R = DexNav, then walk or swim for an unusual signal. Once caught, it stops appearing this way.')
            if method == 'Swarm':
                note(area, 'Swarm: 50% while active; changes hourly. Dates below are days of the month.')

    for den in WIKI.raid_dens():
        area = area_name(den.title.removesuffix(' Raid Den').removeprefix('Underwater '))
        note(area, den.title + ': ' + den.access)
        for name, stars in den.encounters:
            add(area, name, 'Raid den', when=stars, section=den.title,
                detail='Gigantamax' if re.search(r'Gigantamax|Gmax', name, re.I) else '')
            counts['raid rows'] += 1
        counts['dens'] += 1

    book = workbook()
    wild, labels = doc_wild(book)
    differences = []
    for area, place, floor, method, names, colour, cell in wild:
        for name in names:
            species_id = species_ids(name)[0]
            rows = [row for row in by_id[species_id]['w'] if row[0] in by_area[area]['l'] and row[1] != 'Raid den']
            if rows:
                continue
            note_text = 'Doc says %s%s' % (method, ('; ' + place) if place != area else '')
            if method == 'DexNav':
                note_text += '; requires Gen 8 activation'
                note(area, 'Gen 8: register a caught #810-890 species; R = DexNav for unusual signals.')
            add(area, name, method, when=note_text, floor=floor)
            differences.append('%s | %s | %s | %s!%s' % (area, name, note_text, 'Grass/cave' if method == 'Grass/cave' else method, cell))
            counts['doc wild added'] += 1
    for area, method, names, detail in doc_specials(book):
        for name in names:
            species_id = species_ids(name)[0]
            rows = [row for row in by_id[species_id]['w'] if row[0] in by_area[area]['l'] and row[1] != 'Raid den']
            if rows:
                if method == 'Trade' or method == 'Game Corner':
                    for row in rows:
                        if 'Trade' in row[1] or 'Game Corner' in row[1]:
                            row[4] += '; ' + detail if detail not in row[4] else ''
                continue
            add(area, name, method, when='Doc says ' + detail)
            differences.append('%s | %s | Doc says %s: %s' % (area, name, method, detail))
            counts['doc special added'] += 1
    for name, place, schedule in WIKI.swarms():
        area = area_name(place)
        detail = 'Swarm: ' + ', '.join('%s %s' % row for row in schedule)
        species = by_id[species_ids(name)[0]]
        rows = [row for row in species['w'] if row[0] in by_area[area]['l'] and row[1] == 'Swarm']
        if not rows:
            add(area, name, 'Swarm', '50%', when=detail)
        else:
            for row in rows:
                row[4] += '; ' + detail
        note(area, 'Swarm: 50% while active; changes hourly. Dates below are days of the month.')
    sheet = book['Wild Held Items']
    held = []
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row, 3).value or sheet.cell(row, 2).value
        values = [str(sheet.cell(row, column).value).strip() for column in (6, 8) if sheet.cell(row, column).value
                  and str(sheet.cell(row, column).value).strip() != 'X']
        if not name or not values:
            continue
        species_id = sid(NAME_FIX.get(name, name))
        if not species_id:
            raise ValueError('Unresolved held-item species: %s' % name)
        species = by_id[species_id]
        for value in values:
            if not any(key(value) in key(existing) for existing in species.get('hi', [])):
                species.setdefault('hi', []).append(value)
                held.append('%s: %s' % (species['n'], value))
    report('locations: %s' % dict(counts))
    report('doc separator labels: %s' % dict(labels))
    report('DOC SUPPLEMENTS (%d):\n%s' % (len(differences), '\n'.join(differences)))
    report('DOC HELD ADDED (%d):\n%s' % (len(held), '\n'.join(held)))
    differences = membership_differences(data)
    report('RHD MEMBERSHIP DIFFERENCES (%d areas):\n%s' % (len(differences), '\n'.join(
        '%s | RHD only: %s | guide only: %s' % (name, ', '.join(missing) or '-', ', '.join(extra) or '-')
        for name, missing, extra in differences)))
