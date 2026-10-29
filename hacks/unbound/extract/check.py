"""Checks on a built data.json, reproduced from the raw sources (dex part; units 4-7 add theirs).
Usage: uv run python -I hacks/unbound/extract/check.py

- Counts: every RomHackDex page placed (a P entry, merged into one, or dropped for a known reason), M = moves.json less
  the Z / Max copies and the CFRU spellings of one move, AB = every ability a species has.
- Integrity: unique ids and names, valid move / ability indexes, evolutions both ways, every Mega's stone.
- Wiki vs the dex: the number of species whose types / abilities the wiki gives differently (EXPECTED; notes.md lists
  them), so a source update that changes them is noticed.
- Game tables: when hzla's unbound_2_1_constants.js (generated from Skeli's Unbound-Cloud) is under .src/repos/, every
  species' types, stats and abilities against it.
"""
import collections
import glob
import json
import os
from pathlib import Path
import re
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common  # noqa: E402
from common import FORMS, H, RD, WIKI, cfru, key, pid  # noqa: E402

EXPECTED = {'wiki types': 0, 'wiki abilities': 16, 'dropped': 38, 'merged': 116, 'no stone': ['rayquaza_mega'],
            'no description': [], 'not in game table': ['shadow_warrior']}
FAIL = []


def check(ok, msg):
    if not ok:
        FAIL.append(msg)


def main():
    path = os.path.join(H.dir, 'data.json')
    if not os.path.isfile(path):
        print('SKIP unbound extract-check: no data.json')
        return 0
    D = json.load(open(path, encoding='utf-8'))
    M, AB, P = D['M'], D['AB'], D['P']
    S = RD.species()
    ids = {p['id']: p for p in P}

    # ---- every mirror page placed ----
    placed, merged, dropped = 0, 0, collections.Counter()
    for slug, f in FORMS.items():
        if f.kind == 'drop':
            dropped[f.drop] += 1
            continue
        into = cfru.merge_into(f, FORMS, S)
        if into:
            merged += 1
            check(pid(into) in ids, 'merged page %s: target %s not in P' % (slug, pid(into)))
        else:
            placed += 1
            check(pid(slug) in ids, 'mirror page %s not in P' % slug)
    check(placed == len(P), 'P %d entries, mirror pages placed %d' % (len(P), placed))
    check(sum(dropped.values()) == EXPECTED['dropped'], 'dropped %s' % dict(dropped))
    check(merged == EXPECTED['merged'], 'merged %d' % merged)

    # ---- moves and abilities reproduced ----
    raw = json.load(open(os.path.join(H.src, 'romhackdex', 'moves.json'), encoding='utf-8'))
    names = {cfru.official(m['n']) for m in raw if not cfru.move_kind(m['n'])}
    check(set(M) == names and len(M) == len(names), 'M %d, moves.json gives %d official names' % (len(M), len(names)))
    used = {AB[i][0] for p in P for i in p['a'] + ([p['ha']] if p['ha'] >= 0 else [])}
    check(used == {a[0] for a in AB}, 'AB has abilities no species has: %s' % ({a[0] for a in AB} - used))
    nodesc = sorted(a[0] for a in AB if not a[1])
    check(nodesc == EXPECTED['no description'], 'abilities without a description: %s' % nodesc)

    # ---- integrity ----
    for k in ('id', 'n'):
        dup = [x for x, c in collections.Counter(p[k] for p in P).items() if c > 1]
        check(not dup, 'duplicate P %s: %s' % (k, dup))
    nm, na = len(M), len(AB)
    for p in P:
        check(all(0 <= i < nm for _, i in p['lv']) and all(0 <= i < nm for f in ('tm', 'tu', 'eg') for i in p[f]),
              '%s: a move index out of range' % p['id'])
        check(all(0 <= i < na for i in p['a']) and -1 <= p['ha'] < na, '%s: an ability index out of range' % p['id'])
        for how, t in p['et']:
            check(t in ids and [how, p['id']] in ids[t]['ef'], '%s -> %s: no matching ef' % (p['id'], t))
        check(all(t in ids for t in p['of']), '%s: of lists a missing id' % p['id'])
    nostone = [p['id'] for p in P if p['g'] == 1 and not p.get('stone')]
    check(nostone == EXPECTED['no stone'], 'Megas without a stone: %s' % nostone)

    # ---- wiki vs the dex (types, abilities of species the wiki has a page for) ----
    wm = {key(t): m for t, m in WIKI.pokemon().items()}
    abn = [a[0] for a in AB]
    count = collections.Counter()
    for p in P:
        w = wm.get(key(p['n']))
        if not w or p['g'] == 1 or p.get('fc') or common.FORMS_KIND(p['id']) not in ('base', 'regional'):
            continue
        count['wiki types'] += bool(w.types and w.types != p['t'])
        wab = {key(cfru.official(common.WIKI_ABILITY_FIX.get(a, a))) for a, _, _ in w.abilities}
        mine = {key(abn[i]) for i in p['a']} | ({key(abn[p['ha']])} if p['ha'] >= 0 else set())
        count['wiki abilities'] += bool(wab and wab != mine)
    for k in ('wiki types', 'wiki abilities'):
        check(count[k] == EXPECTED[k], '%s: %d species differ (expected %d)' % (k, count[k], EXPECTED[k]))

    # ---- the game's species table (hzla's constants from Unbound-Cloud), when unit 1 has pinned it ----
    found = glob.glob(os.path.join(H.src, 'repos', '**', 'unbound_2_1_constants.js'), recursive=True)
    if found:
        s = open(found[0], encoding='utf-8').read()
        body = s[s.index('return {') + 7:]
        C = json.loads(body[:body.rindex('});')].strip().rstrip(';').strip())
        an, sn = C['ABILITY_DISPLAY_NAMES_BY_DEFINE'], C['SPECIES_DISPLAY_NAMES_BY_DEFINE']
        seen, bad = set(), []
        for define, st in C['BASE_STATS_BY_SPECIES'].items():
            sl = cfru.define_slug(define, FORMS)
            i = (pid(cfru.merge_into(FORMS[sl], FORMS, S) or sl) if sl and FORMS[sl].kind != 'drop' else None) \
                or common.sid(sn.get(define, define))
            if i not in ids or i in seen:
                continue
            seen.add(i)
            p = ids[i]
            t = list(dict.fromkeys(x[5:].capitalize() for x in (st['type1'], st['type2'])))
            s6 = [st[k] for k in ('baseHP', 'baseAttack', 'baseDefense', 'baseSpAttack', 'baseSpDefense', 'baseSpeed')]
            ga = {an.get(st[k], st[k]) for k in ('ability1', 'ability2', 'hiddenAbility') if st.get(k, 'ABILITY_NONE') != 'ABILITY_NONE'}
            mine = {abn[x] for x in p['a']} | ({abn[p['ha']]} if p['ha'] >= 0 else set())
            if t != p['t'] or s6 != p['s'] or ga != mine:
                bad.append(i)
        check(not bad, 'species differing from the game table (%d): %s' % (len(bad), bad[:20]))
        missing = sorted(set(ids) - seen)
        check(missing == EXPECTED['not in game table'], 'P entries not in the game table: %s' % missing)
        print('game table: %d species checked (%s)' % (len(seen), os.path.relpath(found[0], H.dir)))
    else:
        print('game table: not checked (no .src/repos/**/unbound_2_1_constants.js yet; unit 1 pins it)')

    # ---- GID: every species a save can hold reached, every move mapped ----
    G = D.get('GID')
    if G:
        reached = {x for x in G['sp'] if x}
        unreached = sorted(i for i, p in ids.items() if p['g'] != 1 and i not in reached)
        check(unreached == EXPECTED['not in game table'], 'P entries no save id reaches: %s' % unreached)
        check(all(x is None or 0 <= x < nm for x in G['mv']) and sum(1 for x in G['mv'] if x is not None) == len(G['mv']) - 1,
              'GID mv: a move id unmapped or out of range')
        check(all((x is None) == (y is None) for x, y in zip(G['sp'], G['gr'])) and all(x in range(6) for x in G['gr'] if x is not None),
              'GID gr: a species id without a growth rate')
        print('GID: sp %d ids -> %d species, mv %d, it %d' % (sum(1 for x in G['sp'] if x), len(reached),
                                                           sum(1 for x in G['mv'] if x is not None), sum(1 for x in G['it'] if x)))
    else:
        print('GID: not in data.json (no Unbound-Cloud checkout)')

    check_items(D)

    print('dex: P %d (%d mirror pages: %d merged, %d dropped), M %d, AB %d, wiki differences %s'
          % (len(P), len(FORMS), merged, sum(dropped.values()), len(M), len(AB), dict(count)))
    if FAIL:
        print('FAIL unbound extract-check:\n  ' + '\n  '.join(FAIL))
        return 1
    print('OK unbound extract-check')
    return 0


ITEMS_EXPECTED = {'location pages': 67, 'location rows': 687, 'machines': 128, 'stones': 49, 'z': 34, 'tutors': 29,
                  'shop rows': 140, 'missions': 83, 'npcs': 34,
                  # Locations rows whose item no IT row places there (notes.md → Items says why)
                  'not placed': [],
                  # location-guide.xlsx differences (notes.md → Items)
                  'workbook': ['Z-Crystal: workbook only Incinium Z']}     # the workbook: "Currently Unavailable"


def check_items(D):
    """Items (unit 5): counts reproduced from the wiki export, every Locations row placed, GID names resolve."""
    if not D.get('IT'):
        print('items: no IT yet')
        return
    import items
    from wiki import text, item_crystal_parts, useful_npcs
    E = ITEMS_EXPECTED
    box = {b['n']: b['r'] for b in D['IT']}
    by_name = collections.defaultdict(list)
    for b in D['IT']:
        for r in b['r']:
            by_name[r[0]].append((b['n'], r))
    # machines: TM001-TM120 and HM01-HM08, each linked to its move
    tm = box.get('TMs and HMs', [])
    nums = {(m.group(1), int(m.group(2))) for m in (re.match(r'(TM|HM)(\d+) ', r[0]) for r in tm) if m}
    for r in tm:                                    # canon() names a TM by its number ("TM34 (Shockwave)")
        m = re.match(r'(TM|HM)(\d+) ', r[0])
        items.MACHINES[(m.group(1), int(m.group(2)))] = r[0]
    check(len(tm) == E['machines'] and nums == {('TM', i) for i in range(1, 121)} | {('HM', i) for i in range(1, 9)},
          'machines: %d rows, numbers %s' % (len(tm), sorted(nums)[:5]))
    check(all(isinstance(r[2], int) for r in tm), 'machines without a move link: %s'
          % [r[0] for r in tm if not isinstance(r[2], int)])
    # Mega Stones = the dex's stones
    stones = {p['stone']: p['id'] for p in D['P'] if p.get('stone')}
    ms = box.get('Mega Stones', [])
    check(len(ms) == E['stones'] == len(stones) and all(stones.get(r[0]) == r[2] for r in ms),
          'Mega Stones: %d rows, %d P stones' % (len(ms), len(stones)))
    z = box.get('Z-Crystals', [])
    check(len(z) == E['z'] == len(item_crystal_parts(WIKI)), 'Z-Crystals: %d rows' % len(z))
    tut = WIKI.tutors()
    check(len(tut) == E['tutors'] and len(box.get('Move Tutors', [])) == sum(len(t.moves) for t in tut),
          'tutors: %d, rows %d' % (len(tut), len(box.get('Move Tutors', []))))
    check(all(isinstance(r[2], int) for r in box.get('Move Tutors', [])), 'tutor rows without a move link')
    shops = WIKI.item_shops()
    sold = {r[0] for r in box.get('Shops', [])}
    check(len(shops) == E['shop rows'] and all(items.canon(s[2])[0] in sold for s in shops),
          'shops: %d stock rows; not in Shops: %s' % (len(shops), [s[2] for s in shops if items.canon(s[2])[0] not in sold]))
    mis = WIKI.item_missions()
    first = WIKI.page('All Missions')
    check(len(mis) == E['missions'] and first is not None and "A Hero's Journey" in text(first['html']),
          'missions: %d pages (+ #001 in All Missions)' % len(mis))
    npcs = box.get('Useful NPCs', [])
    check(len(npcs) == E['npcs'] == len(useful_npcs(WIKI)), 'Useful NPCs: %d' % len(npcs))
    # Locations pages: every row's place is a page title, every row's item placed there (or by its own page)
    pages = WIKI.item_areas()
    titles = {text(p['title']) for p in pages}
    rows = WIKI.item_locations()
    unmapped = sorted({pl for pl, _, _ in rows} - titles)
    check(len(pages) == E['location pages'] and len(rows) == E['location rows'] and not unmapped,
          'Locations: %d pages, %d rows, places that are no page title: %s' % (len(pages), len(rows), unmapped))
    zn = {r[0] for r in z}
    missing, other = [], []
    for place, raw, _ in rows:
        for one in items.split_names(raw):
            name = items.canon(one)[0]
            m = re.match(r'(TM|HM)\s*0*(\d+)', name)
            if m:
                if (m.group(1), int(m.group(2))) not in nums:
                    other.append('%s: %s (no such machine)' % (place, raw))
            elif name in stones or (re.search(r'ite( [XY])?$', name) and name not in by_name):
                if name not in stones:
                    other.append('%s: %s (not a dex stone)' % (place, raw))
            elif name.endswith(' Z'):
                if name not in zn:
                    other.append('%s: %s (not on the Z-Crystals page)' % (place, raw))
            elif name in items.KEY_DROP:
                pass
            elif not any(place in r[1] for _, r in by_name.get(name, ())):
                missing.append('%s: %s' % (place, raw))
    check(missing == E['not placed'], 'Locations rows no IT row places: %s' % missing)
    check(not other, 'Locations TMs / stones / Z-Crystals not matching their own pages: %s' % other)
    # GID: a game item name that is an IT item is spelled as IT spells it
    bad, hit = [], 0
    for n in (x for x in (D.get('GID') or {}).get('it', []) if x):
        c = items.canon(n)[0]
        if c in by_name:
            hit += 1
            if c != n:
                bad.append('%s -> %s' % (n, c))
    check(not bad, 'GID item names that differ from their IT name: %s' % bad)
    wb = workbook_items(tm, stones, zn)
    check(wb == E['workbook'], 'location-guide.xlsx differences: %s' % wb)
    print('items: %d boxes, %d rows; machines %d, stones %d, Z %d, tutors %d (%d rows), shop stock %d, missions %d, '
          'NPCs %d, Locations %d pages / %d rows (%d places, all page titles), GID items in IT %d'
          % (len(D['IT']), sum(len(b['r']) for b in D['IT']), len(tm), len(ms), len(z), len(tut),
             len(box.get('Move Tutors', [])), len(shops), len(mis), len(npcs), len(pages), len(rows),
             len({p for p, _, _ in rows}), hit))


def workbook_items(tm, stones, zn):
    """location-guide.xlsx sheets TM & HM, Megastone, Z-Crystal against IT (names by key; [] = they agree)."""
    path = os.path.join(H.src, 'research-2026-10-29', 'location-guide.xlsx')
    if not os.path.isfile(path):
        return []
    import openpyxl
    import items
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out, kind, seen = [], 'TM', set()
    mine = {(m.group(1), int(m.group(2))): m.group(3) for m in (re.match(r'(TM|HM)(\d+) (.+)', r[0]) for r in tm)}
    for r in wb['TM & HM'].iter_rows(values_only=True):
        c = [x for x in r if x is not None]
        if c and str(c[0]).strip() == 'No.' and seen:
            kind = 'HM'
        if len(c) < 2 or not str(c[0]).strip().isdigit():
            continue
        k = (kind, int(c[0]))
        seen.add(k)
        theirs = str(c[1]).strip()
        if k in mine and key(cfru.official(theirs)) != key(mine[k]) and key(theirs) != key(mine[k]):
            out.append('%s%s: workbook %s' % (k[0], k[1], theirs))
    out += ['%s%s: not in the workbook' % k for k in sorted(set(mine) - seen)]
    for sheet, have in (('Megastone', set(stones)), ('Z-Crystal', zn)):
        names = set()
        for r in wb[sheet].iter_rows(values_only=True):
            c = [x for x in r if x is not None]
            if len(c) >= 2 and str(c[0]).strip().isdigit():
                names.add(items.canon(str(c[1]).strip())[0])
        out += ['%s: workbook only %s' % (sheet, n) for n in sorted(names - have)]
        out += ['%s: not in the workbook %s' % (sheet, n) for n in sorted(have - names)]
    return out


if __name__ == '__main__':
    sys.exit(main())
