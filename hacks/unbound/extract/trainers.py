"""TR for Pokémon Unbound 2.1.1.1 on Expert: the Expert boss workbook (bosssheet.py), hzla's Expert sets (IVs and the
"hzla says" cross-check), the workbook's cell comments, the wiki's Trainer Locations list (ordinary trainers) with the
ordinary-trainer workbook (held items, scaling, pair splits), and the level caps (trainer-index.xlsx).

    build(D)    writes D['TR'] (boxes = cap splits in walkthrough order, then Post-game); needs dex.build first
    STATS       the report's counts, for check.py

Rules (notes.md → Trainers): the sheet wins; hzla adds IVs (all 31s) on joined slots and its differences as `hn`.
No IVs where hzla has no set (or a set without IVs); no EVs / item / nature where the sheet writes "?" (the raw cell
goes in `x`). Levels are relative to the cap (`r`) where the sheet's comments give the mon's levels at later caps
and every one of them is cap + the same offset; otherwise absolute. Ordinary trainers: the wiki's 309 rows, moves
unknown, placed at the area's earliest visit whose cap covers their levels.
"""
import collections
import copy
import json
import os
import re

import bosssheet
from common import H, WIKI, sid, move, ability, report
from romhackdex import key

STATS = {}

# trainer-index.xlsx → LEVEL CAPS, column D (Difficult / Expert / Insane), rows 2-11.
CAPS = [('Mirskle', 20), ('Vega', 26), ('Alice', 32), ('Mel', 36), ('Maxima', 40), ('Galavan', 45), ('Big Mo', 52),
        ('Tessy', 57), ('Benjamin', 61), ('Elite Four', 75)]
CAP_CELLS = ['LEVEL CAPS!D%d' % r for r in range(2, 12)]
CAP_OF = dict(CAPS, E4=75)
POST = len(CAPS)

# Each split's places in walkthrough order (notes.md → Trainers, placement plan). "#2": the same place met again later
# in the split (its own section).
SECTIONS = [
    ['Frozen Heights', 'Icicle Cave', 'Route 2', 'Route 3', 'Dresco Town'],
    ['Grim Woods', 'Route 4', 'Cinder Volcano', 'Route 5', 'KBT Expressway', 'Crater Town'],
    ['Crater Town', 'Valley Cave', 'Route 6', 'Route 7', 'Frost Mountain', 'Route 8', 'Blizzard City'],
    ['Route 11', 'Cinder Volcano', 'Route 10', 'Route 9', 'Fallshore City'],
    ['Thundercap Mt.', 'Cinder Volcano', 'Tarmigan Mansion'],
    ['Route 12', 'Dehara Department Store', 'Ruins of Void', 'Cinder Volcano', 'Shadow Base', 'Dehara City'],
    ['Route 13', 'Route 14', 'Route 2', 'Auburn Waterway', 'Thundercap Mt.', 'Great Desert', 'Vivill Woods',
     'Vivill Warehouse', 'Distortion World', 'Route 15', 'Route 16', 'Antisis City', 'Antisis Sewers',
     'Antisis City#2'],
    ['Icicle Cave', 'Route 10', 'Route 11', 'Thundercap Mt.', 'Great Desert', 'Route 17', 'Seaport City',
     'Polder Town'],
    ['Cootes Bog', 'Crystal Peak', 'Magnolia Fields', 'Redwood Village'],
    ['Route 18', 'Victory Road', 'Pokémon League'],
    ['Ruins of Void', 'Crystal Peak', 'Redwood Forest', 'Frost Mountain', 'Day Care', 'Blizzard City', 'Title Defense',
     'Battleground'],
]

# Every boss block in walkthrough order: (sheet, block row, split, place, card names, note). A block with several
# names is split by the sheet's "[X's lead]" marks into one card each (fought at once: shared bt).
BOSSES = [
    ('Rival', 2, 0, 'Frozen Heights', ['Rival'], None),
    ('Rival', 21, 0, 'Icicle Cave', ['Rival'], 'Icicle Cave 2F.'),
    ('Gym Leaders', 1, 0, 'Dresco Town', ['Leader Mirskle'], 'Grass gym.'),
    ('Team Shadow Bosses', 2, 1, 'Cinder Volcano', ['Team Shadow Admin Marlon', 'Team Shadow Grunt'],
     'Cinder Volcano 1F, down the ladder right of Chansey.'),
    ('Other Bosses', 41, 1, 'KBT Expressway', ['Black Emboar James'], 'Finale of the Black Emboar quest.'),
    ('Gym Leaders', 21, 1, 'Crater Town', ['Leader Vega'], 'Dark gym.'),
    ('Team Shadow Bosses', 22, 2, 'Crater Town', ['Team Shadow Admin Ivory'], 'After the Crater Gym (kidnapping).'),
    ('Team Shadow Bosses', 42, 2, 'Crater Town', ['Team Shadow Boss Zeph'], 'Right after Ivory.'),
    ('Rival', 40, 2, 'Frost Mountain', ['Rival'], None),
    ('Other Bosses', 61, 2, 'Blizzard City', ['Sinnoh Leader Candice'], None),
    ('Gym Leaders', 41, 2, 'Blizzard City', ['Leader Alice'], 'Flying gym.'),
    ('Other Bosses', 1, 3, 'Cinder Volcano', ['Ace Trainer Alexa'], 'Cinder Volcano 2F.'),
    ('Other Bosses', 81, 3, 'Route 10', ['Terror Granbull Gruff'], 'Finale of the Terror Granbull quest.'),
    ('Team Shadow Bosses', 62, 3, 'Route 9', ['Team Shadow Admin Ivory', 'Team Shadow Grunt'],
     'Route 9 barrier quest finale.'),
    ('Boss Wild Battles', 1, 3, 'Route 9', ['Shadow Warrior'], 'Right after the Ivory tag battle.'),
    ('Gym Leaders', 61, 3, 'Fallshore City', ['Leader Mel'], 'Normal gym.'),
    ('Team Shadow Bosses', 82, 4, 'Thundercap Mt.', ['Team Shadow Admin Marlon'], 'Zapdos capture.'),
    ('Mega Trainers', 41, 4, 'Cinder Volcano', ['Mega Trainer Yale'], 'Cinder Volcano 3F.'),
    ('Gym Leaders', 81, 4, 'Tarmigan Mansion', ['Successor Maxima'], 'Fairy; not a badge.'),
    ('Rival', 59, 5, 'Route 12', ['Rival'], None),
    ('Boss Wild Battles', 20, 5, 'Ruins of Void', ['Claydol'], 'After climbing up a ladder.'),
    ('Boss Wild Battles', 39, 5, 'Ruins of Void', ['Raikou'], 'Three beasts back to back after the sliding puzzle.'),
    ('Boss Wild Battles', 58, 5, 'Ruins of Void', ['Entei'], 'Three beasts back to back after the sliding puzzle.'),
    ('Boss Wild Battles', 77, 5, 'Ruins of Void', ['Suicune'], 'Three beasts back to back after the sliding puzzle.'),
    ('Boss Wild Battles', 96, 5, 'Ruins of Void', ['Ho-Oh'], 'Two birds back to back while meeting Aklove.'),
    ('Boss Wild Battles', 115, 5, 'Ruins of Void', ['Lugia'], 'Two birds back to back while meeting Aklove.'),
    ('Team Light Of Ruin Bosses', 2, 5, 'Cinder Volcano', ['Light of Ruin Agent Vega'], 'After Cube Space.'),
    ('Boss Wild Battles', 134, 5, 'Cinder Volcano', ['Primal Groudon'],
     'After the Vega tag battle; Marlon controls your party\'s attacks.'),
    ('Team Shadow Bosses', 102, 5, 'Shadow Base', ['Team Shadow Boss Zeph'], 'Shadow Base infiltration.'),
    ('Gym Leaders', 101, 5, 'Dehara City', ['Leader Galavan'], 'Electric gym.'),
    ('Mega Trainers', 121, 6, 'Route 13', ['Mega Trainer Osean'], None),
    ('Mega Trainers', 21, 6, 'Route 2', ['Mega Trainer Vivian', 'Mega Trainer Mathis'], None),
    ('Team Light Of Ruin Bosses', 22, 6, 'Vivill Warehouse', ['Light of Ruin Leader Aklove'],
     'Vivill Hideout infiltration.'),
    ('Boss Wild Battles', 153, 6, 'Vivill Warehouse', ['Shadow Warrior'], 'At the entrance; many fights inside.'),
    ('Boss Wild Battles', 172, 6, 'Distortion World', ['Giratina'], 'At the end of the puzzle.'),
    ('Other Bosses', 101, 6, 'Antisis City', ['Black Ferrothorn Rory', 'Black Ferrothorn Chad'], None),
    ('Other Bosses', 121, 6, 'Antisis City', ['Science Society Scientist'], 'Finale of the "Supply and Demand" quest.'),
    ('Other Bosses', 201, 6, 'Antisis City#2', ['Black Ferrothorn Boss Barbs'],
     'Finale of the "Battle of Antisis" quest.'),
    ('Gym Leaders', 121, 6, 'Antisis City#2', ['Leader Big Mo'], 'Fighting gym.'),
    ('Mega Trainers', 1, 7, 'Icicle Cave', ['Mega Trainer Lily'], 'Icicle Cave 1F.'),
    ('Mega Trainers', 61, 7, 'Route 10', ['Mega Trainer Ernesto'], None),
    ('Mega Trainers', 81, 7, 'Route 11', ['Mega Trainer Timothy'], None),
    ('Mega Trainers', 101, 7, 'Thundercap Mt.', ['Mega Trainer Aloysius'], None),
    ('Mega Trainers', 141, 7, 'Great Desert', ['Mega Trainer Travon'], None),
    ('Mega Trainers', 241, 7, 'Route 17', ['Mega Trainer Almuth'], None),
    ('Other Bosses', 161, 7, 'Seaport City', ['Burglar Swizer', 'Black Ferrothorn Bullo'],
     'Finale of the "Home for a Hobo" quest.'),
    ('Other Bosses', 141, 7, 'Seaport City', ['Science Society Scientist'], 'Finale of the "Rogue Electivire" quest.'),
    ('Gym Leaders', 141, 7, 'Polder Town', ['Leader Tessy'], 'Water gym.'),
    ('Rival', 78, 8, 'Cootes Bog', ['Rival'], None),
    ('Boss Wild Battles', 191, 8, 'Crystal Peak', ['Zekrom & Reshiram'], None),
    ('Team Light Of Ruin Bosses', 42, 8, 'Crystal Peak', ['Light of Ruin Admin Ivory'],
     'After the Zekrom / Reshiram fight.'),
    ('Team Light Of Ruin Bosses', 62, 8, 'Crystal Peak', ['Light of Ruin Leader Aklove'],
     'Turning off the Ultimate Weapon.'),
    ('Team Shadow Bosses', 122, 8, 'Crystal Peak', ['Ex Shadow Admin Marlon'], "After Aklove's defeat."),
    ('Boss Wild Battles', 210, 8, 'Crystal Peak', ['Mega Rayquaza & Hoopa'], None),
    ('Gym Leaders', 161, 8, 'Redwood Village', ['Leader Benjamin'], 'Bug gym.'),
    ('Mega Trainers', 161, 9, 'Route 18', ['Mega Trainer Reena'], None),
    ('Mega Trainers', 201, 9, 'Victory Road', ['Mega Trainer Sterling', 'Mega Trainer Portia'], None),
    ('Pokemon League', 1, 9, 'Pokémon League', ['Elite Four Moleman'], None),
    ('Pokemon League', 21, 9, 'Pokémon League', ['Elite Four Elias'], None),
    ('Pokemon League', 41, 9, 'Pokémon League', ['Elite Four Arabella'], None),
    ('Pokemon League', 61, 9, 'Pokémon League', ['Elite Four Penny'], None),
    ('Pokemon League', 81, 9, 'Pokémon League', ['Champion Jax'], None),
    ('Team Light Of Ruin Bosses', 82, 10, 'Ruins of Void', ['Light of Ruin Leader Aklove'],
     'Post-game mission: retrieving the kidnapped mother.'),
    ('Mega Trainers', 181, 10, 'Crystal Peak', ['Mega Trainer Krystal'], None),
    ('Mega Trainers', 221, 10, 'Redwood Forest', ['Mega Trainer Ravonna'], None),
    ('Other Bosses', 181, 10, 'Frost Mountain', ['Expert Ryuzoji', 'Expert Kenlawa'], 'Frost Mountain 2F.'),
    ('Other Bosses', 21, 10, 'Day Care', ['Pokemon Breeder Annie'], 'Day Care final battle.'),
    ('Other Bosses', 221, 10, 'Blizzard City', ['Sinnoh Leader Candice'], 'Post-game rematch.'),
] + [('Title Defense', row, 10, 'Title Defense', [name], None) for row, name in [
    (1, 'Boss Zeph'), (21, 'Admin Marlon'), (41, 'Pokemon Trainer Solenk'), (61, 'Pokemon Trainer Milo'),
    (81, 'Pokemon Trainer Jax'), (101, 'Rival'), (121, 'Rival'), (141, 'Rival')]] + \
    [('Battleground', row, 10, 'Battleground', ['Leader ' + name], None) for row, name in [
        (1, 'Mirskle'), (21, 'Vega'), (41, 'Alice'), (61, 'Mel'), (81, 'Galavan'), (101, 'Big Mo'), (121, 'Tessy'),
        (141, 'Benjamin')]]

# The sheet's export labels for hzla where the label row is empty or differs (key() compares them).
LABEL_FIX = {('Mega Trainers', 241): 'Mega Trainer Almuth', ('Boss Wild Battles', 1): 'Route9',
             ('Boss Wild Battles', 20): 'Ruins of Void', ('Boss Wild Battles', 39): 'Ruins of Void',
             ('Boss Wild Battles', 58): 'Ruins of Void', ('Boss Wild Battles', 77): 'Ruins of Void',
             ('Boss Wild Battles', 96): 'Ruins of Void', ('Boss Wild Battles', 115): 'Ruins of Void',
             ('Boss Wild Battles', 134): 'Cinder Volcano', ('Boss Wild Battles', 153): 'Vivill Hideout',
             ('Boss Wild Battles', 172): 'Distortion World', ('Boss Wild Battles', 191): 'Crystal Peak',
             ('Boss Wild Battles', 210): 'Crystal Peak'}
MOVE_FIX = {'Heatt Wave': 'Heat Wave'}       # Mega Trainers!E205 (recorded in notes.md)
STARTERS = ['Gible', 'Beldum', 'Larvitar']
# The rival partner's "Starter (mega)" (Team Shadow K126 / Light of Ruin K66) by your starter: the rival's Mega.
RIVAL_MEGA = {'Gible': 'Mega Tyranitar', 'Beldum': 'Mega Garchomp', 'Larvitar': 'Mega Metagross'}
MEGA_ABILITY = {'Mega Tyranitar': 'Sand Stream', 'Mega Garchomp': 'Sand Force', 'Mega Metagross': 'Tough Claws'}
BOOST = {'Attack': 0, 'Defense': 1, 'Sp. Atk': 2, 'Sp.Atk': 2, 'Sp. Def': 3, 'Speed': 4}
EV_KEYS = ('hp', 'at', 'df', 'sa', 'sd', 'sp')

# Wiki species typos (the trainer list), by wiki spelling.
WIKI_SPECIES = {'Aurora': 'Aurorus', 'Scrappy': 'Scraggy', 'Wash Rotom': 'Rotom Wash'}
# Wiki area -> the ordinary workbook's sheet.
WB_SHEET = {'Thundercap Mt.': 'Thundercap Mountain', 'Dehara Department Store': 'Dehara City'}
WB_NAME = {'picknicker': 'picnicker'}


def _num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _hzla():
    """hzla's Expert sets: {label key: [(species id, raw species, label, set)]}, and the raw species keys sid can't
    place (malformed: '[2nd', 'Supply and Demand', 'Rogue Electivire')."""
    raw = open(os.path.join(H.src, 'repos', 'hzla', 'backups', 'unbound.js'), encoding='utf-8').read()
    line = next(x for x in raw.splitlines() if x.startswith('formatted_sets["expert"]'))
    sets = json.loads(line.split('=', 1)[1].strip().rstrip(';'))
    out, bad, total = collections.defaultdict(list), [], 0
    for species, labelled in sets.items():
        for label, value in labelled.items():
            total += 1
            pid = sid(species)
            if not pid:
                bad.append((species, label))
                continue
            out[key(re.sub(r'^Lvl \d+ ', '', label))].append((pid, species, label, value))
    return out, bad, total


def _comments():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(H.src, 'trainers-expert.xlsx'))
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.comment:
                    out['%s!%s%d' % (ws.title, c.column_letter, c.row)] = re.sub(r'\s+', ' ', c.comment.text).strip()
    return out


def _caps_check():
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(H.src, 'research-2026-10-29', 'trainer-index.xlsx'),
                                data_only=True)['LEVEL CAPS']
    got = [_num(ws['D%d' % r].value) for r in range(2, 12)]
    assert got == [c for _, c in CAPS], got


def _ladder(text):
    """'Benjamin Cap: 61 Elite Four Cap: 75' -> [('Benjamin', 61), ('Elite Four', 75)]."""
    return [(n.strip(), int(v)) for n, v in re.findall(r'([A-Za-z][A-Za-z ]*?) Cap: (\d+)', text)]


def _ladder_line(ladder):
    parts = ['%d at %s cap' % (lv, 'E4' if n == 'Elite Four' else n) for n, lv in ladder]
    return 'Lv ' + ' · '.join(parts)


def _ordinary_workbook():
    """The ordinary-trainer workbook's blocks: {sheet: [block]}, block = dict(row, names, double, mons, items)."""
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(H.src, 'research-2026-10-29', 'ordinary-trainers.xlsx'), data_only=True)
    out = {}
    for ws in wb.worksheets:
        blocks = []
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value != 'Name':
                continue
            title = bosssheet.clean(ws.cell(r, 2).value)
            double = '[Double Battle]' in title
            parts = [p.strip() for p in re.sub(r'\s*\[[^\]]*\]', '', title).split(' & ')]
            first = parts[0].split()
            names = [parts[0]] + [p if len(p.split()) > 1 else ' '.join(first[:-1] + [p]) for p in parts[1:]]
            items = bosssheet.clean(ws.cell(r + 1, 2).value)
            mons, owner = [], 0
            for c in range(2, 30):
                sp = bosssheet.clean(ws.cell(r + 3, c).value)
                if not sp:
                    continue
                lead = re.search(r"\[(.+?)'s lead\]", sp)
                if lead:
                    owner = next(i for i, n in enumerate(names) if n.split()[-1] == lead[1])
                mons.append(dict(sp=re.sub(r'\s*\[[^\]]*\]', '', sp), min=_num(ws.cell(r + 4, c).value),
                                 item=bosssheet.clean(ws.cell(r + 5, c).value), owner=owner,
                                 cell='%s!%s%d' % (ws.title, ws.cell(r + 3, c).column_letter, r + 3)))
            blocks.append(dict(sheet=ws.title, row=r, title=title, names=names, double=double, mons=mons,
                               items=items, used=False))
        out[ws.title] = blocks
    return out


def _wb_sid(name):
    n = re.sub(r'^A\. ', 'Alolan ', re.sub(r'^G\. ', 'Galarian ', name)).replace('Ferrothon', 'Ferrothorn')
    return sid(n)


def _nkey(name):
    k = key(name)
    for a, b in WB_NAME.items():
        k = k.replace(a, b)
    return k


def build(D):
    _caps_check()
    move_index = {m: i for i, m in enumerate(D['M'])}
    ab_index = {a[0]: i for i, a in enumerate(D['AB'])}
    pnames = {p['id']: p['n'] for p in D['P']}
    hz, hz_bad, hz_total = _hzla()
    comments = _comments()
    used_comments = set()
    blocks = {(b.sheet, b.row): b for b in bosssheet.blocks()}
    assert len(blocks) == 84 and sum(len(g.mons) for b in blocks.values() for g in b.groups) == 439
    counts = collections.Counter()
    differs, sheet_only, unknown, relative, absolute_ladder, fixes = [], [], [], [], [], []
    joined_sets = set()

    def moves_of(raw, cell, x):
        out, hp = [], None
        for m in raw:
            if m == 'default moveset':
                x.append('default moveset')
                fixes.append('%s: "default moveset", no moves listed' % cell)
                continue
            m = MOVE_FIX.get(m, m)
            h = re.fullmatch(r'Hidden Power (\w+)', m)
            if h:
                hp, m = h[1], 'Hidden Power'
            name = move(m)
            assert name, (cell, m)
            out.append(move_index[name])
        return out, hp

    def ab(name, cell):
        a = ability(name)
        assert a, (cell, name)
        return ab_index[a]

    def mon_of(slot, block, hzset=None):
        """One sheet slot -> mon (without hzla fields), its raw ability text."""
        x = []
        raw = slot.species
        ld = re.search(r'\[(1st|2nd) lead\]', raw)
        pre = '(pre-burned)' in raw
        if '(respawnable)' in raw:
            x.append('respawnable')
        name = re.sub(r'\s*\[[^\]]*\]|\s*\((pre-burned|respawnable)\)', '', raw).strip()
        pid = sid(name)
        assert pid, (slot.cell, raw)
        mon = dict(id=pid, l=slot.level)
        # ability: "A / B" random slot, "A → B" before / after Mega Evolution
        text = slot.ability
        before, _, after = text.partition(' → ')
        if after:
            mon['a'] = ab(after, slot.cell)
            if ' / ' in before:
                mon['ao'] = text
            else:
                mon['a0'] = ab(before, slot.cell)
        elif ' / ' in text:
            options = text.split(' / ')
            pick = next((o for o in options if hzset and key(o) == key(hzset.get('ability', ''))), options[0])
            mon['a'] = ab(pick, slot.cell)
            mon['ao'] = text
        else:
            mon['a'] = ab(text, slot.cell)
        item = slot.item
        if item.endswith('?'):
            x.append('item ' + item)
        elif item:
            mon['i'] = item
        if slot.nature.endswith('?'):
            x.append('nature ' + slot.nature)
        elif slot.nature:
            mon['na'] = slot.nature
        evs = [_num(v) for v in slot.evs]
        if all(v is not None for v in evs):
            mon['ev'] = evs
        else:
            x.append('EVs ' + ' / '.join(slot.evs))
        mon['m'], hp = moves_of(slot.moves, slot.cell, x)
        if hp:
            mon['hp'] = hp
        if ld:
            mon['ld'] = 1 if ld[1] == '1st' else 2
        if pre:
            mon['pre'] = ['Burned']
        if block.sheet == 'Boss Wild Battles' and slot.label not in ('none', ''):
            boosts = slot.label
            if boosts == 'none?':
                x.append('boosts none?')
            else:
                au = [0] * 5
                for part in [p.strip() for p in boosts.split(',')]:
                    m = re.fullmatch(r'(?:([A-Za-z. ]+?) ([+-]\d)|([+-]\d) ([A-Za-z. ]+))', part)
                    if part == 'Raid Shields':
                        x.append('Raid Shields')
                    elif part.startswith('Omniboost'):
                        au = [int(part.split()[-1])] * 5
                    elif m:
                        stat, v = (m[1], m[2]) if m[1] else (m[4], m[3])
                        au[BOOST[stat.strip()]] = int(v)
                    else:
                        raise AssertionError((slot.cell, part))
                if any(au):
                    mon['au'] = au
        # comments on the slot's cells: species (evolution by level), ability, moves
        for r, what in [(slot.row, 'species'), (slot.row + 2, 'ability')] + \
                [(slot.row + 4 + i, 'move') for i in range(4)]:
            cell = '%s!%s%d' % (slot.sheet, slot.col, r)
            if cell in comments:
                used_comments.add(cell)
                c = comments[cell]
                x.append(c if what == 'species' else '%s: %s' % (bosssheet.clean(_cell_value(block, slot, r)), c)
                         if what == 'move' else c)
        if x:
            mon['x'] = x
        return mon

    def _cell_value(block, slot, r):
        return {slot.row + 4 + i: m for i, m in enumerate(slot.moves)}.get(r, '')

    def hz_label(block, slot):
        label = LABEL_FIX.get((block.sheet, block.row))
        if block.sheet == 'Rival' and block.row == 40:
            label = 'Rival 3 |Player Chose %s|' % next(g.starter for g in block.groups if slot in g.mons)
        if label is None:
            label = re.sub(r'\s*(- Expert|\[Expert\])$', '', slot.label).strip()
        return key(label)

    def compare(mon, slot, hzset):
        """The hzla set's differences from the sheet, as "hzla says" items."""
        out = []
        if hzset.get('level') != slot.level:
            out.append('level %s' % hzset.get('level'))
        text = slot.ability
        options = {key(o) for o in re.split(r' / | → ', text)}
        if key(hzset.get('ability', '')) not in options:
            out.append('ability %s' % (hzset.get('ability') or 'none'))
        hzi = hzset.get('item', '') or ''
        if key(hzi) != key(slot.item.rstrip('?')):
            out.append('item %s' % (hzi or 'none'))
        theirs = sorted(key(MOVE_FIX.get(m, m)) for m in hzset.get('moves', []) if m)
        ours = sorted(key(MOVE_FIX.get(m, m)) for m in slot.moves if m != 'default moveset')
        if theirs != ours:
            out.append('moves ' + ', '.join(m for m in hzset.get('moves', []) if m))
        if key(hzset.get('nature', '')) != key(slot.nature.rstrip('?')):
            out.append('nature %s' % hzset.get('nature'))
        ev = hzset.get('evs') or {}
        typo = 'def' in ev
        hzev = [ev.get(k, ev.get('def', 0) if k == 'df' else 0) for k in EV_KEYS]
        sheet_ev = [_num(v.rstrip('?')) for v in slot.evs]
        if None not in sheet_ev and hzev != sheet_ev:
            out.append('EVs ' + (' / '.join(str(v) for v in hzev) if any(hzev) else 'none'))
        return out, typo

    def join(mon, slot, block, hzlabel):
        counts['slots'] += 1
        cands = [c for c in hz.get(hzlabel, []) if c[0] == mon['id'] or
                 (sid(re.sub(r'-Shield$', '', c[1])) == mon['id'])]
        if not cands:
            mon['src'] = 's'
            sheet_only.append(slot.cell)
            return None
        pid, species, label, hzset = cands[0]
        joined_sets.add((species, label))
        counts['joined'] += 1
        if hzset.get('ivs'):
            mon['iv'] = [hzset['ivs'].get(k, 31) for k in EV_KEYS]
            assert set(mon['iv']) == {31}, slot.cell
        out, typo = compare(mon, slot, hzset)
        if typo:
            fixes.append("%s: hzla's EV key 'def' read as Defense" % slot.cell)
        if out:
            differs.append((slot.cell, out))
            mon['hn'] = out
        return hzset

    def team(block, group, hzlabel_of=None):
        out = []
        for slot in group.mons:
            if slot.species == 'Starter (mega)':
                out.append(None)
                continue
            label = hz_label(block, slot)
            cands = [c for c in hz.get(label, []) if c[0] == sid(re.sub(r'\s*\[[^\]]*\]|\s*\(.*?\)', '',
                                                                      slot.species).strip())]
            mon = mon_of(slot, block, cands[0][3] if cands else None)
            join(mon, slot, block, label)
            if 'iv' not in mon:
                unknown.append((slot.cell, 'IVs'))
            if 'ev' not in mon:
                unknown.append((slot.cell, 'EVs'))
            mon['_slot'] = slot
            out.append(mon)
        return out

    def levels(t, mons, block, cap):
        """Relative levels (r) where the comments give the mon's levels at later caps, all at cap + one offset."""
        for mon in mons:
            slot = mon['_slot']
            cell = '%s!%s%d' % (slot.sheet, slot.col, slot.row + 1)
            if cell not in comments:
                continue
            used_comments.add(cell)
            ladder = _ladder(comments[cell])
            assert ladder, cell
            mon.setdefault('x', []).insert(0, _ladder_line(ladder))
            offsets = {lv - CAP_OF[n] for n, lv in ladder}
            named = re.fullmatch(r'Level at (.+) Cap', block.levelrow)
            base = CAP_OF[named[1]] if named and named[1] in CAP_OF else cap
            if len(offsets) == 1 and base and slot.level - base in offsets:
                mon['r'] = slot.level - base
                del mon['l']
                relative.append(slot.cell)
            elif len(offsets) == 1 and named and named[1] in CAP_OF:
                # The block names its cap but its level isn't that cap + the notes' offset (Reena: Lv 59-61 "at E4
                # Cap", notes 75-77 there). The notes follow every other Mega Trainer's pattern: they win.
                (r,) = offsets
                mon['x'].insert(1, 'Sheet says Lv %d' % slot.level)
                mon['r'] = r
                del mon['l']
                relative.append(slot.cell)
                fixes.append('%s: Lv %d at %s Cap, the notes say %d: relative to the notes' % (cell, slot.level, named[1], base + r))
            else:
                absolute_ladder.append((slot.cell, slot.level, base, sorted(offsets)))
                note = 'Levels as the sheet lists them: its notes give other levels at this cap (see each Pokémon).'
                if note not in t['note']:
                    t['note'].append(note)

    # boxes
    groups = [dict(n='%s Split [Cap %d]' % (name, cap), cap=cap, tab='', sec=[]) for name, cap in CAPS]
    groups.append(dict(n='Post-game', cap=None, tab='', sec=[]))
    groups[0]['note'] = ['Expert teams. Boss levels are at the level cap unless the trainer says "minimum levels".']
    sections = copy.deepcopy(SECTIONS)
    secs = [{h: dict(h=h.split('#')[0], tr=[]) for h in hs} for hs in sections]
    battle = [0]
    wiki_rows = WIKI.trainer_locations()
    boss_names = collections.Counter()

    def place(split, h, trainer):
        if h not in secs[split]:
            secs[split][h] = dict(h=h.split('#')[0], tr=[])
            sections[split].insert(len(sections[split]) - 1, h)
            report('  trainers: place %s added to %s' % (h, groups[split]['n']))
        if split < POST:
            trainer['cap'] = CAPS[split][1]
        secs[split][h]['tr'].append(trainer)

    def notes_of(block, own):
        notes = [own] if own else []
        head = block.head
        cap_word = block.levelrow
        if cap_word == 'Minimum Level':
            notes.append('Minimum levels (the sheet).')
        elif 'UNSURE' in cap_word:
            notes.append('Levels at the cap; the sheet marks them unsure.')
        name = head.get('name', '')
        if 'NOT UPDATED' in name:
            notes.append('The sheet says this team is not updated to 2.1.0.')
        if head.get('items'):
            notes.append('Items: ' + head['items'] + '.')
        for gcell, comment in comments.items():
            sheet, rest = gcell.split('!')
            row = int(re.sub(r'\D', '', rest))
            first = min(slot.row for g in block.groups for slot in g.mons)
            if sheet == block.sheet and rest.startswith('B') and block.row <= row < first and head.get('gim'):
                notes.append(comment)       # the Fight Gimmicks cell's explanation
                used_comments.add(gcell)
        return notes

    def tags_of(block):
        tags = []
        gim = block.head.get('gim', '')
        for part in [p.strip() for p in re.split(r',\s*', gim) if p.strip()]:
            if part.lower() == 'negative energy':
                part = 'Negative Energy'
            tags.append(part)
        for m in re.findall(r'\[([^\]]+Battle)\]', block.head.get('name', '')):
            if m not in tags:
                tags.append(m)
        if block.sheet == 'Boss Wild Battles':
            tags.insert(0, 'Wild')
        return tags

    rival5_mega = {}
    for sheet, row, split, h, names, own in BOSSES:
        block = blocks.pop((sheet, row))
        cap = CAPS[split][1] if split < POST else None
        foes = [g for g in block.groups if g.role == 'foe']
        partners = [g for g in block.groups if g.role == 'partner']
        tags, notes = tags_of(block), notes_of(block, own)
        if sheet == 'Boss Wild Battles' and block.side:
            for title, items in block.side.items():
                notes.append('%s: %s.' % (title.capitalize(), ', '.join(items)))
        pts = []
        for g in partners:
            mons = team(block, g)
            pname = g.partner.replace('Pokemon Trainer ', 'Pokémon Trainer ')
            if None in mons:        # the rival's Mega depends on your starter
                i = mons.index(None)
                slot = g.mons[i]
                comment = comments.get('%s!%s%d' % (slot.sheet, slot.col, slot.row + 4))
                used_comments.add('%s!%s%d' % (slot.sheet, slot.col, slot.row + 4))
                for st in STARTERS:
                    mega = RIVAL_MEGA[st]
                    ref = rival5_mega[st]
                    found = re.search(re.escape(mega) + r' moveset: - (.+?)(?= Mega |$)', comment)
                    mv = [m.strip() for m in found[1].split(' - ')]
                    mon = dict(id=sid(mega), l=slot.level, a=ab(MEGA_ABILITY[mega], slot.cell), i=ref['i'],
                               na=ref['na'], ev=list(ref['ev']), m=[move_index[move(m)] for m in mv], src='s',
                               x=['Nature and EVs from Rival 5\'s %s (the sheet: "check corresponding spread")' % mega])
                    team_st = [dict((k, v) for k, v in m.items() if k != '_slot') for m in mons[:i]] + [mon] + \
                        [dict((k, v) for k, v in m.items() if k != '_slot') for m in mons[i + 1:]]
                    pts.append(dict(n=pname, st=st, p=team_st))
                counts['slots'] += 1
                sheet_only.append(slot.cell)
                unknown.append((slot.cell, 'IVs'))
            else:
                pts.append(dict(n=pname, p=[dict((k, v) for k, v in m.items() if k != '_slot') for m in mons]))
        cards = []
        for g in foes:
            mons = team(block, g)
            if sheet == 'Rival' and block.row == 78:
                rival5_mega[g.starter] = mons[-1]
            # split a pair's team by its "[X's lead]" marks
            if len(names) > 1:
                owners, owner = [], -1
                for m in mons:
                    lead = re.search(r"\[(.+?)'s lead\]", m['_slot'].species)
                    if lead:
                        owner = next(i for i, n in enumerate(names) if n.split()[-1].startswith(lead[1][:4]))
                    owners.append(owner)
                assert -1 not in owners and len(set(owners)) == len(names), (sheet, row, owners)
                for i, n in enumerate(names):
                    cards.append((n, g.starter, [m for m, o in zip(mons, owners) if o == i]))
            else:
                cards.append((names[0], g.starter, mons))
        st_of = {101: 'Gible', 121: 'Beldum', 141: 'Larvitar'} if sheet == 'Title Defense' else {}
        bt = None
        if len(names) > 1:
            battle[0] += 1
            bt = 'unbound-%d' % battle[0]
        for n, starter, mons in cards:
            t = dict(k='', n=n, loc=h.split('#')[0], tags=list(tags), note=list(notes), p=mons)
            if len(names) > 1 and 'Double Battle' not in t['tags'] and not any(x.startswith('Tag Battle')
                                                                                 for x in t['tags']):
                t['tags'].append('Double Battle')
            st = starter or st_of.get(row)
            if st:
                t['st'] = st
            if pts:
                t['pt'] = copy.deepcopy(pts)
            if bt:
                t['bt'] = bt
            levels(t, mons, block, cap)
            for m in mons:
                m.pop('_slot', None)
            if sheet == 'Other Bosses' and row == 201:
                t['note'].append('Its minimum levels are above the Big Mo cap.')
            place(split, h, t)
            if sheet == 'Mega Trainers' or (sheet, row) in (('Other Bosses', 1), ('Other Bosses', 181)):
                boss_names[key(n)] += 1     # the wiki lists these with the ordinary trainers: the same card
        counts['boss blocks'] += 1
    assert not blocks, list(blocks)
    for cell in comments:
        if cell not in used_comments:
            report('  trainers: comment not used: %s' % cell)
    counts['comments used'] = len(used_comments)

    # ordinary trainers: the wiki list
    wb = _ordinary_workbook()
    wb_all = [b for bs in wb.values() for b in bs]
    ordinary, boss_rows, name_only = [], [], []
    wiki_by_name = {}
    rows = list(wiki_rows)
    for i, row in enumerate(rows):
        if key(row.name) in boss_names or (row.shared and key(row.shared) in boss_names):
            boss_rows.append(row)
            continue
        ordinary.append(row)
    wb_matched, wb_items, loose = 0, 0, []

    def wb_find(row):
        sheet = WB_SHEET.get(row.area, row.area)
        for b in wb.get(sheet, []):
            if _nkey(row.name) in [_nkey(n) for n in b['names']]:
                return b
        for b in wb_all:
            if _nkey(row.name) in [_nkey(n) for n in b['names']]:
                loose.append('%s (%s) -> %s!B%d' % (row.name, row.area, b['sheet'], b['row']))
                return b
        return None

    def wiki_sid(name, block_mons):
        n = WIKI_SPECIES.get(name, name)
        if n == 'Nidoran':
            got = [m['sp'] for m in block_mons if m['sp'].startswith('Nidoran ')]
            return (sid(got[0]), None) if got else (sid('Nidoran F'), 'Nidoran: ♀ or ♂ not given (wiki, workbook)')
        pid = sid(n)
        assert pid, name
        return pid, None

    combined = {}       # first row name -> its wiki team, for name-only rows
    cards = []
    for row in ordinary:
        b = wb_find(row)
        idx = b['names'].index(next(n for n in b['names'] if _nkey(n) == _nkey(row.name))) if b else None
        if b:
            b['used'] = True
            wb_matched += 1
        wmons = row.mons
        if row.shared:
            wmons = combined.get(row.shared, [])
        mons, notes, own = [], [], None
        pair = bool(b) and len(b['names']) > 1
        bmons = b['mons'] if b else []
        if pair:
            own = [m for m in bmons if m['owner'] == idx]
        # wteam: [(species id, wiki level or the workbook minimum, workbook mon or None, warning)]
        if row.shared or (pair and len(wmons) == len(bmons) and len(wmons) > len(own)):
            # the wiki lists the pair's combined team under the first name: keep this trainer's part
            if not row.shared:
                combined[row.name] = wmons
            if pair:
                pool = [(wiki_sid(n, bmons), lv) for n, lv in wmons]
                wteam = []
                for m in own:
                    p = _wb_sid(m['sp'])
                    hit = next((w for w in pool if w[0][0] == p), None)
                    if hit:
                        pool.remove(hit)
                        wteam.append((p, hit[1], m, hit[0][1]))
                    else:
                        wteam.append((p, m['min'], m, 'Level: the workbook minimum (the wiki has no %s)' % m['sp']))
                    assert p, m
            elif b:
                # the name-only row's own workbook block: its species, levels from the combined wiki team
                pool = [(wiki_sid(n, bmons), lv) for n, lv in wmons]
                wteam = []
                for m in bmons:
                    p = _wb_sid(m['sp'])
                    hit = next((w for w in pool if w[0][0] == p), None)
                    if hit:
                        pool.remove(hit)
                    wteam.append((p, hit[1] if hit else m['min'], m, hit[0][1] if hit else
                                  'Level: the workbook minimum (the wiki has no %s)' % m['sp']))
            else:
                wteam = None
        elif b and not pair and len(bmons) < len(wmons) and any(r.shared == row.name for r in ordinary):
            # the wiki's combined team for this row and the name-only one after it: keep the workbook's part
            combined[row.name] = wmons
            pool, wteam = list(bmons), []
            for n, lv in wmons:
                pid, warn = wiki_sid(n, bmons)
                hit = next((m for m in pool if _wb_sid(m['sp']) == pid), None)
                if hit:
                    pool.remove(hit)
                    wteam.append((pid, lv, hit, warn))
        else:
            wteam = [wiki_sid(n, bmons) + (lv,) for n, lv in wmons]
            wteam = [(pid, lv, None, warn) for pid, warn, lv in wteam]
            if b:
                bm = own if pair else bmons
                if len(bm) == len(wteam):
                    wteam = [(pid, lv, bm[i], warn) for i, (pid, lv, _, warn) in enumerate(wteam)]
                else:
                    pool, out = list(bm), []
                    for pid, lv, _, warn in wteam:
                        hit = next((m for m in pool if _wb_sid(m['sp']) == pid), None)
                        if hit:
                            pool.remove(hit)
                        out.append((pid, lv, hit, warn))
                    wteam = out
                    report('  ordinary: %s (%s): wiki %d Pokémon, workbook %d' % (row.name, row.area, len(wmons),
                                                                                  len(bm)))
        if wteam is None:
            name_only.append(row.name)
            notes.append('Team unknown: the wiki lists one team for %s and %s.' % (row.shared, row.name))
            wteam = []
        for pid, lv, m, warn in wteam:
            mon = dict(id=pid, l=lv, m=[])
            x = [warn] if warn else []
            if m:
                if m['item']:
                    mon['i'] = m['item']
                    wb_items += 1
                wsp = _wb_sid(m['sp'])
                if wsp != pid or m['min'] != lv:
                    x.append('Workbook: %s, minimum Lv %s' % (pnames.get(wsp, m['sp']), m['min']))
            if x:
                mon['x'] = x
            mons.append(mon)
        if b:
            mins = [m['min'] for m in (own if pair else b['mons']) if m['min']]
            if mins:
                notes.append('Scales with your level, evolving as it does: minimum Lv %s (workbook).' %
                             ('%d' % min(mins) if min(mins) == max(mins) else '%d-%d' % (min(mins), max(mins))))
            if b['items']:
                notes.append('Items: %s.' % b['items'])
        if row.place:
            notes.insert(0, row.place)
        t = dict(k='', n=row.name, loc=row.area, tags=['Moves unknown'], note=notes, p=mons)
        if pair:
            t['bt'] = 'unbound-w%d-%s' % (b['row'], key(b['sheet']))
            if b['double']:
                t['tags'].append('Double Battle')
            mates = [n for n in b['names'] if _nkey(n) != _nkey(row.name)]
            t['note'].append('Fought together with %s.' % ' and '.join(mates))
        cards.append((row, t))

    # place: the area's earliest visit whose cap covers the team's levels
    for row, t in cards:
        top = max([m['l'] for m in t['p']] or [0])
        area = row.area
        visits = [s for s in range(POST + 1) if area in SECTIONS[s] and (s == POST or CAPS[s][1] >= top)]
        if visits:
            split = visits[0]
        else:
            split = next(s for s in range(POST + 1) if s == POST or CAPS[s][1] >= top)
            report('  ordinary: %s (%s, Lv %d): no visit to the area covers it; placed in %s' %
                   (row.name, area, top, groups[split]['n']))
            counts['ordinary placed by cap only'] += 1
        place(split, area, t)
        # an ordinary trainer goes before the bosses of its place
        sec = secs[split][area]['tr']
        sec.insert(next((i for i, x in enumerate(sec) if 'Moves unknown' not in x['tags']), len(sec) - 1), sec.pop())

    for s, hs in enumerate(sections):
        for h in hs:
            if secs[s][h]['tr']:
                groups[s]['sec'].append(secs[s][h])
    for gi, g in enumerate(groups):
        for si, s in enumerate(g['sec']):
            for ti, t in enumerate(s['tr']):
                t['k'] = '%d.%d.%d' % (gi, si, ti)
    D['TR'] = dict(g=groups, caps=[[n if n != 'Maxima' else 'Successor Maxima', c] for n, c in CAPS])

    wb_only = [b for b in wb_all if not b['used']]
    STATS.update(dict(slots=counts['slots'], joined=counts['joined'], sets=len(joined_sets), differs=differs,
                      sheet_only=sheet_only, hz_bad=hz_bad, hz_total=hz_total, unknown=unknown, relative=relative,
                      absolute_ladder=absolute_ladder, wiki=len(rows), ordinary=len(ordinary), boss_rows=len(boss_rows),
                      name_only=name_only, wb_matched=wb_matched, wb_items=wb_items, wb_only=wb_only, loose=loose,
                      comments=len(comments), comments_used=len(used_comments), fixes=fixes))
    report('trainers: %d sheet slots in %d boss blocks; hzla %d sets: %d slots joined (%d sets), %d differ, %d sheet '
           'only, %d malformed hzla keys' % (counts['slots'], counts['boss blocks'], hz_total, counts['joined'],
                                             len(joined_sets), len(differs), len(sheet_only), len(hz_bad)))
    report('trainers: wiki %d rows = %d ordinary + %d boss cards; %d name-only rows with no team; workbook %d '
           'matched, %d held items, %d workbook-only blocks' % (len(rows), len(ordinary), len(boss_rows),
                                                              len(name_only), wb_matched, wb_items, len(wb_only)))
    report('trainers: %d mons at relative levels; %d with a later-cap line kept absolute; %d comments used of %d' %
           (len(relative), len(absolute_ladder), len(used_comments), len(comments)))
    for cell, out in differs:
        report('  hzla differs: %s: %s' % (cell, '; '.join(out)))
    for cell in sheet_only:
        report('  sheet only: %s' % cell)
    for species, label in hz_bad:
        report('  hzla malformed key: %r / %r' % (species, label))
    for cell, what in unknown:
        report('  unknown %s: %s' % (what, cell))
    for item in absolute_ladder:
        report('  later-cap levels not cap + one offset (absolute): %s Lv %d, cap %s, offsets %s' % item)
    for item in fixes:
        report('  sheet fix: %s' % item)
    for item in loose:
        report('  workbook loose match: %s' % item)
    for b in wb_only:
        report('  workbook only: %s!B%d %s' % (b['sheet'], b['row'], b['title']))
