import openpyxl,json,re
key=lambda s:re.sub(r'[^a-z0-9]','',str(s).lower())
D=json.load(open('data.json')); R='/home/claude/ei/miriamlefae_eidex/eidex-HEAD/src/data/'
MD={key(m['name']):m for m in json.load(open(R+'moveData.json')).values()}
AD={key(a['name']):a for a in json.load(open(R+'abilityData.json')).values()}
TN={int(k):v['typeName'] for k,v in json.load(open(R+'typeData.json')).items()}
CAT=['PHYSICAL','SPECIAL','STATUS']
wb=openpyxl.load_workbook('/mnt/user-data/uploads/1_3_1_Emerald_Imperium_Changes__Stats__Types__Abilities__Moves_.xlsx',data_only=True)
mv={};conf=[];miss=[]
for r in list(wb['Move Changes'].iter_rows(values_only=True))[1:]:
    if not r[1]: continue
    nm=re.sub(r'\(NEW!\)','',str(r[1])).strip(); k=key(nm)
    eff=(r[9] or '').strip(); mv[k]={'n':nm,'e':eff,'new':'NEW' in str(r[1])}
    m=MD.get(k)
    if not m: miss.append(nm); continue
    d=[]
    if key(r[3])!=key(TN[m['type']]): d.append(f"type doc {r[3]} dex {TN[m['type']]}")
    if r[5]!=CAT[m['cat']]: d.append(f"category doc {r[5]} dex {CAT[m['cat']]}")
    bp=None if r[6] in ('—','-',None) else r[6]; dbp=m['power'] or None
    if bp!=dbp and not (dbp==1): d.append(f"power doc {r[6]} dex {m['power']}")
    ac=None if r[7] in ('—','-',None) else r[7]; dac=None if m['acc'] in (0,999) else m['acc']
    if ac!=dac: d.append(f"accuracy doc {r[7]} dex {m['acc']}")
    if r[8]!=m['pp']: d.append(f"PP doc {r[8]} dex {m['pp']}")
    if d: conf.append((nm,d))
print('moves in doc',len(mv),'not in dex',miss); print('CONFLICTS',len(conf))
for c in conf: print(c)
ab={};amiss=[]
for r in list(wb['Ability Changes'].iter_rows(values_only=True))[1:]:
    if not r[1]: continue
    nm=re.sub(r'\(NEW!\)','',str(r[1])).strip(); ab[key(nm)]={'n':nm,'e':r[2].strip(),'new':'NEW' in str(r[1])}
    if key(nm) not in AD: amiss.append(nm)
print('abilities in doc',len(ab),'not in dex data',amiss)
# moves the dex flagged changed vs Gen9 but the doc does not mention
Mk=[key(n) for n in D['M']]
und=[D['M'][i] for i,v in enumerate(D['MV']) if (v['h'] or v.get('x')) and Mk[i] not in mv]
print('dex-detected changes not in doc:',len(und),und)
json.dump({'mv':mv,'conf':conf,'miss':miss,'ab':ab,'amiss':amiss,'undoc':und},open('ck/changes_moves.json','w'))
