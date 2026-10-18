# Stage 4: Items tab from the official "Item, TMs, Mega Stones etc and Useful NPC Locations" xlsx.
# D.IT = [{n section, k kind, r:[[item, text, link]], g?}]  link: move index (TM/HM) or mega Pokemon id.
import openpyxl,json,re
key=lambda s:re.sub(r'[^a-z0-9]','',str(s).lower())
D=json.load(open('data.json')); P=D['P']; Mk={key(m):i for i,m in enumerate(D['M'])}
wb=openpyxl.load_workbook('/mnt/user-data/uploads/Item__TMs__Mega_Stones__etc__and_Useful_NPC_Locations_.xlsx',data_only=True)
V=lambda x:'' if x is None else str(x).strip()
log=[];IT=[]
def mv(s):
    m=re.match(r'^(TM|HM)\s*0*(\d+)\s*-\s*(.+)$',s.strip())
    if not m: return s,None
    n=m[3].strip(); n={'Drain Kiss':'Draining Kiss'}.get(n,n); i=Mk.get(key(n))
    if i is None: log.append('TM/HM move not in dex: '+s)
    return f'{m[1]}{int(m[2]):03d} {n}' if m[1]=='TM' else f'{m[1]}{int(m[2]):02d} {n}',i
rows=list(wb['TM & HM Locations'].iter_rows(values_only=True))[2:]
IT.append({'n':'TMs','k':'tm','r':[[*mv(V(r[0]))[:1],V(r[1]),mv(V(r[0]))[1]] for r in rows if V(r[0])]})
IT.append({'n':'HMs','k':'tm','r':[[*mv(V(r[5]))[:1],V(r[6]),mv(V(r[5]))[1]] for r in rows if len(r)>5 and V(r[5])]})
# Mega Stones: pairs in columns 1-2 and 5-6
megas=[p for p in P if p['g']==1]
def stone2mega(st):
    s=key(st); best=None;bl=0
    for p in megas:
        b=key(re.sub(r'^(Mega|Primal) ','',p['n']).split(' ')[0]); suf=p['n'].split(' ')[-1] if p['n'].count(' ')>=2 else ''
        l=0
        while l<min(len(s),len(b)) and s[l]==b[l]: l+=1
        if l>=3 and (l>bl or (l==bl and suf and s.endswith(key(suf)))):
            if suf and not s.endswith(key(suf)) and re.search(r'(x|y|o|d)$',s): continue
            best,bl=p,l
    return best
mr=[]
for r in list(wb['Mega Stones'].iter_rows(values_only=True))[1:]:
    for a,b in ((1,2),(5,6)):
        if len(r)>b and V(r[a]) and V(r[a])!='Mega Stone': mr.append([V(r[a]),V(r[b]),None])
for x in mr:
    p=stone2mega(x[0])
    if p: x[2]=p['id']; p['stone']=x[0]; p['sloc']=x[1]
    else: log.append('Mega Stone with no Mega form in the dex: '+x[0])
IT.append({'n':'Mega Stones','k':'mega','r':mr})
nost=[p['n'] for p in megas if 'stone' not in p and not p['n'].startswith('Primal')]
if nost: log.append('Megas with no stone location in the sheet: '+', '.join(nost))
IT.append({'n':'Battle item rewards','k':'gen','r':[[V(r[0]),V(r[1]),None] for r in list(wb['Battle Item Rewards'].iter_rows(values_only=True))[1:] if V(r[0])]})
# Marts and NPCs: header rows (single cell) start groups
def grouped(ws,col):
    out=[];g=None
    for r in ws.iter_rows(values_only=True):
        c=[V(x) for x in r]
        if not any(c): continue
        cells=[x for x in c[col:] if x]
        if len(cells)==1 and g is not None and not g['r'] and not cells[0].startswith('-'):
            g['h']+=' · '+cells[0]; continue
        if len(cells)==1 and not cells[0].startswith('-') and not re.match(r'^(TM|HM)\s*\d',cells[0]) and not (g and g.get('items')):
            g={'h':cells[0],'r':[]};out.append(g)
        elif len(cells)==1 and not cells[0].startswith('-') and g and len(cells[0])>40:
            g['h']+=' ('+cells[0]+')'
        else:
            if g is None: g={'h':'','r':[]};out.append(g)
            g['r'].append(cells)
    return out
marts=[]
for g in grouped(wb['Poké Marts items'],1):
    for c in g['r']:
        price=c[1] if len(c)>1 else ''
        price='Free' if price in ('Free','0.0') else (f'₽{int(float(price))}' if re.match(r'^\d+(\.\d+)?$',price) else price)
        marts.append([c[0].lstrip('*'),g['h']+(' · '+price if price else '')+(' · needs Ability Patch unlocked' if c[0].startswith('*') else ''),None])
        if len(c)>2: marts.append([c[2],g['h']+' · Free',None])
IT.append({'n':'Shops','k':'gen','r':marts,'note':'Any Poké Mart sells the vanilla stock unless listed here.'})
npc=[]
for g in grouped(wb['Useful NPC Locations'],1):
    for c in g['r']:
        for part in re.split(r'\n\s*-\s*',c[0]): npc.append([g['h'],part.lstrip('- ').strip(),None])
IT.append({'n':'Useful NPCs','k':'npc','r':npc})
for s in IT:
    for r in s['r']:
        if re.search(r'missable|miss (it|this) |only once|before (you|defeating)|permanent',r[1],re.I): r.append(1)
D['IT']=IT; D['src']+='; Items xlsx (EI)'
json.dump(D,open('data.json','w'),separators=(',',':'),ensure_ascii=False)
print({s['n']:len(s['r']) for s in IT}); print('\n'.join(log))
print('applite',[ (r[0],r[2]) for r in mr if 'ppli' in r[0]]); print('flagged',[ (r[0],r[1][:60]) for s in IT for r in s['r'] if len(r)>3][:12])
print(marts[:4],npc[:3])
json.dump(log,open('ck/items_log.json','w'))
